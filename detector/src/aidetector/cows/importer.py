import argparse
import csv
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from aidetector.cows.registry import (
    CowRegistry,
    NumberTaken,
    normalize_life_number,
    normalize_number,
)

# Column names as they appear in exports of herd management programs, such as
# Lely's: Diernr, Resp 1, Levensnummer, Gesl, Naam, Werknummer, Diercat.
_LIFE_COLUMNS = ("levensnummer", "levnr", "life number", "lifenumber", "i&r")
# "Levnr moeder" is the mother's life number.
_PARENT_WORDS = ("moeder", "vader", "mother", "father")
# The farmer calls a cow by her collar number (Lely: Diernr); a heifer without
# a collar by her work number.
_COLLAR_COLUMNS = ("halsband", "diernr", "diernummer", "collar")
_WORK_COLUMNS = ("werknummer", "werknr")
_OTHER_NUMBER_COLUMNS = ("nummer", "number")
# The responder hangs on the collar: without one, a heifer has no collar yet.
_RESPONDER_COLUMNS = ("resp", "transponder")
_NAME_COLUMNS = ("naam", "name", "roepnaam")
_SEX_COLUMNS = ("gesl", "sex", "gender")
_CATEGORY_COLUMNS = ("diercat", "categorie", "category")
# Only cows and heifers take part. With a list of categories (see
# cows.herd_categories) only those count; without one, calves and bulls are
# recognised by these words and anything unclear stays in.
_MALE_VALUES = ("m", "man", "mannelijk", "male", "stier", "bull")
_SKIPPED_CATEGORIES = ("kalf", "kalveren", "calf", "stier", "bull")
DEFAULT_CATEGORIES = ("Koeien", "Vrouwelijk jongvee")
# Exports often start with a title or the farm's name above the header.
_HEADER_ROWS = 10
EXCEL_SUFFIXES = (".xlsx", ".xlsm")
# More missing animals than this share of the herd is more likely a partial
# export than animals that left, so nobody is archived then.
_MAX_MISSING = 0.2


def _column(
    header: list[str], names: tuple[str, ...], skip: set[int], avoid: tuple[str, ...] = ()
) -> int | None:
    lowered = [column.strip().lower() for column in header]
    for name in names:
        for index, column in enumerate(lowered):
            if (
                index not in skip
                and name in column
                and not any(word in column for word in avoid)
            ):
                return index
    return None


def _cell(row: list, index: int | None) -> str:
    return _text(row[index]) if index is not None and index < len(row) else ""


def _text(value: object) -> str:
    """Excel gives numbers as floats: 30.0 is collar number 30."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_table(path: Path) -> list[list[str]]:
    """Reads the first sheet of an Excel file, or a CSV with ; , or tab."""
    if path.suffix.lower() in EXCEL_SUFFIXES:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            return [[_text(value) for value in row] for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    text = path.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    if not lines:
        return []
    dialect = csv.Sniffer().sniff(lines[0], delimiters=";,\t")
    return [[cell.strip() for cell in row] for row in csv.reader(lines, dialect)]


@dataclass
class HerdRow:
    line: int
    number: str
    life_number: str
    name: str | None
    # Whether the number is a collar number, or else a work number.
    collar: bool


@dataclass
class ParsedHerd:
    animals: list[HerdRow] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)
    # Animals of other categories, such as calves and male animals.
    skipped: int = 0


@dataclass
class _Columns:
    life: int
    collar: int | None = None
    work: int | None = None
    other: int | None = None
    responder: int | None = None
    name: int | None = None
    sex: int | None = None
    category: int | None = None

    @classmethod
    def find(cls, header: list[str]) -> "_Columns | None":
        life = _column(header, _LIFE_COLUMNS, set(), _PARENT_WORDS)
        if life is None:
            return None
        taken = {life}

        def take(names: tuple[str, ...]) -> int | None:
            index = _column(header, names, taken, _PARENT_WORDS)
            if index is not None:
                taken.add(index)
            return index

        # The responder first, so "Resp 1" is not read as a number column.
        responder = take(_RESPONDER_COLUMNS)
        columns = cls(
            life=life,
            responder=responder,
            collar=take(_COLLAR_COLUMNS),
            work=take(_WORK_COLUMNS),
            other=take(_OTHER_NUMBER_COLUMNS),
        )
        if columns.collar is None and columns.work is None and columns.other is None:
            return None
        columns.name = take(_NAME_COLUMNS)
        columns.sex = take(_SEX_COLUMNS)
        columns.category = take(_CATEGORY_COLUMNS)
        return columns

    def number(self, row: list[str]) -> tuple[str, bool]:
        """Her collar number when she has a collar, else her work number."""
        collar, work, other = _cell(row, self.collar), _cell(row, self.work), _cell(row, self.other)
        has_collar = bool(collar) and (self.responder is None or bool(_cell(row, self.responder)))
        if has_collar:
            return collar, True
        if work:
            return work, False
        if other:
            return other, self.work is None and self.collar is None
        return collar, True

    def skip(self, row: list[str], categories: Sequence[str] | None) -> bool:
        if _cell(row, self.sex).lower() in _MALE_VALUES:
            return True
        if self.category is None:
            return False
        category = _cell(row, self.category).casefold()
        if categories is not None:
            return category not in {wanted.strip().casefold() for wanted in categories}
        return any(word in category for word in _SKIPPED_CATEGORIES)


def parse_herd(
    rows: list[list[str]], categories: Sequence[str] | None = DEFAULT_CATEGORIES
) -> ParsedHerd:
    """Finds the header and returns the cows and heifers with valid numbers,
    a problem per row that could not be read, and how many other animals
    (calves, male animals) were left out."""
    columns, data = None, list(enumerate(rows, start=1))
    for index, row in enumerate(rows[:_HEADER_ROWS]):
        columns = _Columns.find(row)
        if columns is not None:
            data = data[index + 1 :]
            break
    if columns is None:
        # No recognised header: number, life number, name.
        columns = _Columns(life=1, collar=0, name=2)

    parsed = ParsedHerd()
    for line, row in data:
        if not any(_text(cell) for cell in row):
            continue
        if columns.skip(row, categories):
            parsed.skipped += 1
            continue
        number, collar = columns.number(row)
        try:
            parsed.animals.append(
                HerdRow(
                    line,
                    normalize_number(number),
                    normalize_life_number(_cell(row, columns.life)),
                    _cell(row, columns.name) or None,
                    collar,
                )
            )
        except ValueError as error:
            parsed.problems.append(f"Regel {line}: {error}")
    return parsed


@dataclass
class ImportResult:
    added: int = 0
    problems: list[str] = field(default_factory=list)
    # How many animals got their collar number, and how many their work number.
    by_collar: int = 0
    by_work: int = 0
    skipped: int = 0

    def summary(self) -> str:
        """The first line of the reply to the farmer."""
        animals = "dier" if self.added == 1 else "dieren"
        if self.by_collar and self.by_work:
            text = (
                f"✅ {self.added} {animals} ingelezen: {self.by_collar} op halsbandnummer, "
                f"{self.by_work} op werknummer (pinken zonder halsband)."
            )
        elif self.by_work:
            text = f"✅ {self.added} {animals} ingelezen op werknummer."
        else:
            text = f"✅ {self.added} {'koe' if self.added == 1 else 'koeien'} ingelezen."
        if self.skipped:
            text += f" {self.skipped} andere dieren (kalveren, mannelijk) overgeslagen."
        return text


def import_cows(
    path: Path, registry: CowRegistry, categories: Sequence[str] | None = DEFAULT_CATEGORIES
) -> ImportResult:
    """Adds each animal of the file (collar or work number, life number,
    optional name); animals already known keep what they have unless the
    file gives them a free number."""
    rows = read_table(path)
    if not rows:
        return ImportResult(problems=["Het bestand is leeg"])
    parsed = parse_herd(rows, categories)
    result = ImportResult(problems=parsed.problems, skipped=parsed.skipped)
    for animal in parsed.animals:
        try:
            registry.add(animal.number, animal.life_number, animal.name)
        except NumberTaken as error:
            result.problems.append(
                f"Regel {animal.line}: nummer {error.number} hoort al bij {error.holder},"
                " gebruik /wissel in Telegram"
            )
            continue
        result.added += 1
        if animal.collar:
            result.by_collar += 1
        else:
            result.by_work += 1
    result.problems.sort(key=lambda problem: int(problem.split()[1].rstrip(":")))
    return result


@dataclass
class SyncResult:
    added: list[str] = field(default_factory=list)
    renumbered: list[str] = field(default_factory=list)
    archived: list[str] = field(default_factory=list)
    returned: list[str] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)
    skipped: int = 0

    @property
    def changed(self) -> bool:
        return bool(self.added or self.renumbered or self.archived or self.returned)

    def summary(self) -> str:
        parts = []
        if self.added:
            parts.append(f"{len(self.added)} nieuw")
        if self.returned:
            parts.append(f"{len(self.returned)} terug")
        if self.renumbered:
            parts.append(f"{len(self.renumbered)} ander nummer")
        if self.archived:
            parts.append(f"{len(self.archived)} weg (archief)")
        return "📋 Koeienlijst bijgewerkt: " + (", ".join(parts) or "niets veranderd") + "."


def sync_herd(
    path: Path,
    registry: CowRegistry,
    at: datetime | None = None,
    categories: Sequence[str] | None = DEFAULT_CATEGORIES,
) -> SyncResult:
    """Makes the register follow the herd list, which is leading: new animals
    are added, numbers and names follow the list, and animals that are no
    longer on it are archived. Numbers change in two passes, so two cows can
    swap collars."""
    at = at or datetime.now()
    parsed = parse_herd(read_table(path), categories)
    result = SyncResult(problems=parsed.problems, skipped=parsed.skipped)
    listed: dict[str, HerdRow] = {}
    numbers: dict[str, str] = {}
    for animal in parsed.animals:
        if animal.life_number in listed:
            result.problems.append(f"Regel {animal.line}: {animal.life_number} staat er twee keer in")
        elif animal.number in numbers:
            result.problems.append(
                f"Regel {animal.line}: nummer {animal.number} staat ook bij {numbers[animal.number]}"
            )
        else:
            listed[animal.life_number] = animal
            numbers[animal.number] = animal.life_number
    if not listed:
        result.problems.append("Geen dieren gevonden in de lijst; ik verander niets.")
        return result

    active = {cow.life_number for cow in registry.active_cows()}
    missing = sorted(active - set(listed))
    archive = missing
    if active and len(missing) > max(3, _MAX_MISSING * len(active)):
        result.problems.append(
            f"{len(missing)} van de {len(active)} dieren staan niet in de lijst. Dat lijkt "
            "een halve export, dus ik archiveer niemand. Klopt het wel? Gebruik dan /weg."
        )
        archive = []
    for life_number in archive:
        result.archived.append(registry.label(life_number))
        registry.archive(life_number, at)

    # First free the numbers that go to another animal, then hand them out.
    for life_number, animal in listed.items():
        current = registry.number_of(life_number, at)
        if current is not None and current != animal.number:
            registry.unassign(life_number, at)
    for life_number, animal in listed.items():
        cow = registry.cow(life_number)
        known = cow is not None and cow.archived is None
        current = registry.number_of(life_number, at) if known else None
        try:
            registry.add(animal.number, life_number, animal.name, at=at)
        except NumberTaken as error:
            result.problems.append(
                f"Regel {animal.line}: nummer {animal.number} hoort nog bij {error.holder}"
            )
            continue
        label = registry.label(life_number, at)
        if cow is None:
            result.added.append(label)
        elif not known:
            result.returned.append(label)
        elif current != animal.number:
            result.renumbered.append(label)
    return result


def _configured_directory() -> Path:
    from aidetector.utils.config import ChatConfig, config

    for detector in config.detectors:
        telegram = detector.exporters.telegram if detector.exporters else None
        for chat in [telegram] if isinstance(telegram, ChatConfig) else telegram or []:
            if chat.cows:
                return (
                    chat.cows.directory
                    or chat.feedback_directory / "koeien"
                ).expanduser()
    raise SystemExit(
        "Geen telegram.cows in config.json; geef de map op met --directory"
    )


def main() -> None:
    parser = argparse.ArgumentParser(
        prog="import-koeien",
        description="Leest koeien en pinken uit een Excel- of CSV-bestand met halsband-"
        " of werknummer, levensnummer en (optioneel) naam, bijvoorbeeld een export van"
        " het managementprogramma.",
    )
    parser.add_argument("bestand", type=Path)
    parser.add_argument(
        "--directory",
        type=Path,
        help="Map met de koeien; standaard die uit config.json",
    )
    arguments = parser.parse_args()
    directory = arguments.directory or _configured_directory()
    result = import_cows(arguments.bestand, CowRegistry(directory.resolve()))
    print(f"{result.summary().removeprefix('✅ ')} Map: {directory}")
    for problem in result.problems:
        print(problem)
