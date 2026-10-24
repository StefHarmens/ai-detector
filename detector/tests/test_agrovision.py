import os
import stat

import pytest
from openpyxl import Workbook

from aidetector.cows.agrovision import ConnectorError, get_connector

HEADER = ["Diernr", "Resp 1", "Levensnummer", "Gesl", "Geb dat", "Naam", "Werknummer", "Diercat", "Levnr moeder", "Aankdat"]


def herd(path, count):
    workbook = Workbook()
    workbook.active.append(HEADER)
    for number in range(1, count + 1):
        workbook.active.append([number, None, f"NL{number:09d}", "V", None, f"Koe{number}", number, "Koeien", None, None])
    workbook.save(path)


LOGIN = {"url": "https://agrovision.example/rapport", "username": "boer", "password": "geheim"}


@pytest.fixture
def connector(tmp_path):
    connector = get_connector(tmp_path)
    connector.fetch = lambda state, out: herd(out, 10)
    return connector


def test_save_keeps_the_password_private(connector):
    connector.save(LOGIN)
    view = connector.view()
    assert "password" not in view and view["has_password"] and view["configured"]
    assert stat.S_IMODE(os.stat(connector.folder / "connector.json").st_mode) == 0o600
    # Empty password on a later save keeps the old one.
    connector.save({**LOGIN, "password": "", "run_at": "04:00"})
    assert connector._read()["password"] == "geheim"
    assert connector.view()["run_at"] == "04:00"


@pytest.mark.parametrize("change", [{"url": "agrovision"}, {"username": ""}, {"run_at": "25:00"}])
def test_save_rejects_incomplete_input(connector, change):
    with pytest.raises(ValueError):
        connector.save({**LOGIN, **change})


def test_menu_path_is_stored_as_steps(connector):
    connector.save({**LOGIN, "path": "Dieranalyse >  Koeienlijst > "})
    assert connector._read()["steps"] == ["Dieranalyse", "Koeienlijst"]
    assert connector.view()["path"] == "Dieranalyse > Koeienlijst"
    connector.save({**LOGIN, "path": ""})
    assert connector._read()["steps"] == []


def test_success_writes_the_herd_list_and_status(connector):
    connector.save(LOGIN)
    connector._run()
    view = connector.view()
    assert view["status"] == "ok" and view["animals"] == 10 and view["succeeded"]
    assert connector.herd_file.is_file()


def test_failure_keeps_the_old_list_and_gives_the_reason(connector):
    connector.save(LOGIN)
    connector._run()
    told = []
    connector.notify = told.append

    def refuse(state, out):
        raise ConnectorError("Inloggen mislukt")

    connector.fetch = refuse
    before = connector.herd_file.read_bytes()
    connector._run()
    connector._run()
    view = connector.view()
    assert view["status"] == "failed" and view["reason"] == "Inloggen mislukt"
    assert view["animals"] == 10
    assert connector.herd_file.read_bytes() == before
    # Told once per streak, and again when it works.
    assert len(told) == 1
    connector.fetch = lambda state, out: herd(out, 10)
    connector._run()
    assert len(told) == 2 and connector.view()["status"] == "ok"


def test_partial_list_and_non_list_are_refused(connector):
    connector.save(LOGIN)
    connector._run()
    connector.fetch = lambda state, out: herd(out, 5)
    connector._run()
    assert "onvolledig" in connector.view()["reason"]
    connector.fetch = lambda state, out: out.write_text("<html>Inloggen</html>")
    connector._run()
    assert connector.view()["status"] == "failed"
    assert len(__import__("aidetector.cows.importer", fromlist=["x"]).parse_herd(
        __import__("aidetector.cows.importer", fromlist=["x"]).read_table(connector.herd_file)).animals) == 10


def test_delete_forgets_everything(connector):
    connector.save(LOGIN)
    connector.delete()
    assert not connector.view()["configured"]


def test_session_code_is_taken_out_of_the_address(connector):
    connector.save(
        {
            **LOGIN,
            "url": "https://cowvision.agrovision.nl/CowVision/Applications/Portal_v2026_03/?SessionID=abc#analysis/reports/5/293/1/stallijst%20hoentjen",
        }
    )
    assert connector.view()["url"] == (
        "https://cowvision.agrovision.nl/CowVision/Applications/Portal_v2026_03/"
        "#analysis/reports/5/293/1/stallijst%20hoentjen"
    )


def test_one_list_file_is_kept_and_replaced(connector):
    connector.save(LOGIN)
    connector._run()
    connector.fetch = lambda state, out: herd(out, 12)
    connector._run()
    connector.fetch = lambda state, out: out.write_text("<html>")
    connector._run()
    assert sorted(path.name for path in connector.directory.glob("*.xlsx*")) == ["koeienlijst.xlsx"]
    assert not list(connector.folder.glob("*.xlsx*"))
    assert connector.view()["animals"] == 12


XML = """<?xml version="1.0"?>
<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">
<Worksheet ss:Name="Stallijst"><Table>
<Row><Cell><Data ss:Type="String">Stallijst</Data></Cell></Row>
<Row>{header}</Row>
{rows}
</Table></Worksheet></Workbook>"""


def cell(value, kind="String"):
    return f'<Cell><Data ss:Type="{kind}">{value}</Data></Cell>'


def xml_list(count):
    header = "".join(cell(name) for name in HEADER)
    rows = "\n".join(
        "<Row>"
        + cell(number, "Number")
        + '<Cell ss:Index="3"><Data ss:Type="String">NL%09d</Data></Cell>' % number
        + cell("V")
        + cell("")
        + cell(f"Koe{number}")
        + cell(number, "Number")
        + cell("Koeien")
        + "</Row>"
        for number in range(1, count + 1)
    )
    return XML.format(header=header, rows=rows)


def test_excel_2003_xml_download_is_converted(connector):
    connector.save(LOGIN)
    connector.fetch = lambda state, out: out.write_text(xml_list(10))
    connector._run()
    assert connector.view()["status"] == "ok" and connector.view()["animals"] == 10
    assert connector.herd_file.read_bytes()[:2] == b"PK"


def test_html_table_download_is_converted(connector):
    connector.save(LOGIN)
    rows = "".join(
        f"<tr><td>{n}</td><td></td><td>NL{n:09d}</td><td>V</td><td></td><td>Koe{n}</td><td>{n}</td><td>Koeien</td></tr>"
        for n in range(1, 4)
    )
    head = "<tr>" + "".join(f"<th>{name}</th>" for name in HEADER[:8]) + "</tr>"
    connector.fetch = lambda state, out: out.write_text(f"<html><table>{head}{rows}</table></html>")
    connector._run()
    assert connector.view()["animals"] == 3


def test_unknown_file_type_names_how_it_starts(connector):
    connector.save(LOGIN)
    connector.fetch = lambda state, out: out.write_text("%PDF-1.4 rest")
    connector._run()
    assert "PDF" in connector.view()["reason"]


REAL_SHAPE = """<?xml version="1.0"?>
<?mso-application progid="Excel.Sheet"?>
<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet"
 xmlns:o="urn:schemas-microsoft-com:office:office"
 xmlns:x="urn:schemas-microsoft-com:office:excel"
 xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">
 <Worksheet ss:Name="stallijst">
  <Table ss:ExpandedColumnCount="10">
   <Row>
    <Cell><Data ss:Type="String">Diernr</Data><NamedCell ss:Name="_FilterDatabase"/></Cell>
    <Cell><Data ss:Type="String">Resp 1</Data><NamedCell ss:Name="_FilterDatabase"/></Cell>
    <Cell><Data ss:Type="String">Levensnummer</Data><NamedCell
      ss:Name="_FilterDatabase"/></Cell>
    <Cell><Data ss:Type="String">Gesl</Data></Cell>
    <Cell><Data ss:Type="String">Geb dat</Data></Cell>
    <Cell><Data ss:Type="String">Naam</Data></Cell>
    <Cell><Data ss:Type="String">Haarkl</Data></Cell>
    <Cell><Data ss:Type="String">Diercat</Data></Cell>
    <Cell><Data ss:Type="String">Levnr moeder</Data></Cell>
    <Cell><Data ss:Type="String">Aankdat</Data></Cell>
   </Row>
   <Row>
    <Cell><Data ss:Type="String">7</Data></Cell>
    <Cell><Data ss:Type="String">1234567</Data></Cell>
    <Cell ss:StyleID="s62"><Data ss:Type="String">NL 0000.0003.0</Data></Cell>
    <Cell><Data ss:Type="String">Vrouwelijk</Data></Cell>
    <Cell><Data ss:Type="String">19-03-20</Data></Cell>
    <Cell><Data ss:Type="String">BERTHA 7</Data></Cell>
    <Cell><Data ss:Type="String">Roodbont</Data></Cell>
    <Cell><Data ss:Type="String">Koeien</Data></Cell>
    <Cell><Data ss:Type="String">NL 000000001</Data></Cell>
    <Cell><Data ss:Type="String"></Data></Cell>
   </Row>
   <Row>
    <Cell ss:Index="3" ss:StyleID="s62"/>
   </Row>
  </Table>
 </Worksheet>
</Workbook>"""


def test_the_shape_of_agrovisions_own_export(connector):
    connector.save(LOGIN)
    connector.fetch = lambda state, out: out.write_text(REAL_SHAPE)
    connector._run()
    assert connector.view()["status"] == "ok" and connector.view()["animals"] == 1


def test_a_list_without_animals_says_what_it_holds(connector):
    connector.save(LOGIN)
    connector.fetch = lambda state, out: herd(out, 0)
    connector._run()
    reason = connector.view()["reason"]
    assert "1 gevulde rijen" in reason and "Diernr" in reason


def test_a_refused_download_is_kept_until_the_next_good_one(connector):
    connector.save(LOGIN)
    connector.fetch = lambda state, out: herd(out, 0)
    connector._run()
    assert (connector.folder / "laatste-download.xml").is_file()
    connector.fetch = lambda state, out: herd(out, 10)
    connector._run()
    assert not (connector.folder / "laatste-download.xml").exists()


BOLD_SHAPE = """<?xml version="1.0" encoding="UTF-8"?>
<?mso-application progid="Excel.Sheet"?>
<Workbook xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet" xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:x="urn:schemas-microsoft-com:office:excel" xmlns="urn:schemas-microsoft-com:office:spreadsheet">
<Styles><Style ss:ID="s61"><NumberFormat ss:Format="General Date"/></Style></Styles>
<Worksheet ss:Name="Stallijst"><Table>
<Column ss:Width="100"></Column>
<Row>
<Cell ss:Index="1"><Data ss:Type="String">Diernr</Data></Cell><Cell ss:Index="2"><Data ss:Type="String">Resp 1</Data></Cell><Cell ss:Index="3"><Data ss:Type="String">Levensnummer</Data></Cell><Cell ss:Index="4"><Data ss:Type="String">Gesl</Data></Cell><Cell ss:Index="5"><Data ss:Type="String">Geb dat</Data></Cell><Cell ss:Index="6"><Data ss:Type="String">Naam</Data></Cell><Cell ss:Index="7"><Data ss:Type="String">Haarkl</Data></Cell><Cell ss:Index="8"><Data ss:Type="String">Diercat</Data></Cell><Cell ss:Index="9"><Data ss:Type="String">Levnr moeder</Data></Cell><Cell ss:Index="10"><Data ss:Type="String">Aankdat</Data></Cell></Row>
<Row>
<Cell ss:Index="1" ><Data ss:Type="String">1</Data></Cell><Cell ss:Index="2" ><Data ss:Type="String">1234567</Data></Cell><Cell ss:Index="3" ><Data ss:Type="String">NL 0000.<b>0003</b>.0</Data></Cell><Cell ss:Index="4" ><Data ss:Type="String">Vrouwelijk</Data></Cell><Cell ss:Index="5" ><Data ss:Type="String">19-03-20</Data></Cell><Cell ss:Index="6" ><Data ss:Type="String">BERTHA 1</Data></Cell><Cell ss:Index="7" ><Data ss:Type="String">Roodbont</Data></Cell><Cell ss:Index="8" ><Data ss:Type="String">Koeien</Data></Cell><Cell ss:Index="9" ><Data ss:Type="String">NL 000000001</Data></Cell><Cell ss:Index="10" ><Data ss:Type="String"></Data></Cell></Row>
<Row>
<Cell ss:Index="1" ><Data ss:Type="String">2</Data></Cell><Cell ss:Index="2" ><Data ss:Type="String">1234568</Data></Cell><Cell ss:Index="3" ><Data ss:Type="String">NL 0000.<b>0007</b>.7</Data></Cell><Cell ss:Index="4" ><Data ss:Type="String">Vrouwelijk</Data></Cell><Cell ss:Index="5" ><Data ss:Type="String">18-04-23</Data></Cell><Cell ss:Index="6" ><Data ss:Type="String">KLAARTJE 2</Data></Cell><Cell ss:Index="7" ><Data ss:Type="String">Roodbont</Data></Cell><Cell ss:Index="8" ><Data ss:Type="String">Koeien</Data></Cell><Cell ss:Index="9" ><Data ss:Type="String">NL 000000002</Data></Cell><Cell ss:Index="10" ><Data ss:Type="String"></Data></Cell></Row>
</Table></Worksheet></Workbook>"""


def test_life_number_with_bold_part_is_read_whole(connector):
    from openpyxl import load_workbook

    connector.save(LOGIN)
    connector.fetch = lambda state, out: out.write_text(BOLD_SHAPE)
    connector._run()
    assert connector.view()["status"] == "ok" and connector.view()["animals"] == 2
    sheet = load_workbook(connector.herd_file).active
    assert [row[2] for row in sheet.iter_rows(min_row=2, values_only=True)] == ["NL 0000.0003.0", "NL 0000.0007.7"]


def sheet_file(path, rows):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)


COW_REPORT = [["Diernr", "Status", "Dgn status", "Lactdgn", "#Insem"], ["8", "Drachtig", "112", "187", "2"]]
HEIFER_REPORT = [["Diernr", "Lft mdn", "Status", "Dgn status", "#Insem"], ["8835", "16", "Tochtig", "1", "0"]]
REPORTS = {
    "report_koeien": "https://cowvision.agrovision.nl/CowVision/Applications/Portal_v2026_03/?SessionID=abc#analysis/reports/5/301/1/vruchtbaarheid%20mk",
    "report_jongvee": "Mijn rapporten > Vruchtbaarheid JV",
}


def fetch_with(connector, cows=COW_REPORT, heifers=HEIFER_REPORT):
    def fetch(state, out):
        herd(out, 10)
        if cows is not None:
            sheet_file(connector.folder / "download-koeien.xml", cows)
        if heifers is not None:
            sheet_file(connector.folder / "download-jongvee.xml", heifers)

    connector.fetch = fetch


def test_reports_are_a_route_or_a_menu_path_and_lose_their_session_code(connector):
    connector.save({**LOGIN, **REPORTS})

    extras = connector._extras(connector._read())

    assert [(extra["key"], extra["fragment"], extra["steps"]) for extra in extras] == [
        ("koeien", "analysis/reports/5/301/1/vruchtbaarheid%20mk", []),
        ("jongvee", "", ["Mijn rapporten", "Vruchtbaarheid JV"]),
    ]
    assert "SessionID" not in connector.view()["reports"]["koeien"]["value"]
    # Without reports nothing extra is fetched.
    connector.save(LOGIN)
    assert connector._extras(connector._read()) == []


def test_a_report_address_needs_a_route(connector):
    with pytest.raises(ValueError, match="route na de #"):
        connector.save({**LOGIN, "report_koeien": "https://cowvision.agrovision.nl/portal/"})


def test_fertility_is_fetched_with_the_herd_list_and_kept_per_animal(connector):
    from aidetector.cows import fertility

    connector.save({**LOGIN, **REPORTS})
    fetch_with(connector)

    connector._run()

    view = connector.view()
    assert view["status"] == "ok"
    assert view["reports"]["koeien"]["status"] == "ok" and view["reports"]["koeien"]["animals"] == 1
    assert view["reports"]["jongvee"]["status"] == "ok"
    data = fertility.load(connector.directory)
    assert fertility.lookup(data, "8")["lact_days"] == 187
    assert fertility.lookup(data, "8835")["age_months"] == 16
    assert not list(connector.folder.glob("download-*"))


def test_a_failing_fertility_report_leaves_the_herd_list_alone(connector):
    from aidetector.cows import fertility

    connector.save({**LOGIN, **REPORTS})
    # The cows' report lacks a column; the heifers' report was not downloaded.
    fetch_with(connector, cows=[["Diernr", "Status"], ["8", "Drachtig"]], heifers=None)
    connector.extra_reasons = {}

    def fetch(state, out):
        herd(out, 10)
        sheet_file(connector.folder / "download-koeien.xml", [["Diernr", "Status"], ["8", "Drachtig"]])
        connector.extra_reasons = {"jongvee": "De knop 'exporteren naar Excel' staat niet op de pagina"}

    connector.fetch = fetch
    connector._run()

    view = connector.view()
    assert view["status"] == "ok" and view["animals"] == 10
    assert view["reports"]["koeien"]["status"] == "failed"
    assert "Dgn status" in view["reports"]["koeien"]["reason"] and "Lactdgn" in view["reports"]["koeien"]["reason"]
    assert view["reports"]["jongvee"]["status"] == "failed"
    assert "exporteren" in view["reports"]["jongvee"]["reason"]
    assert (connector.folder / "laatste-download-koeien.xml").is_file()
    assert fertility.load(connector.directory) == {}

    # Fixed: fetched again, the kept download goes.
    fetch_with(connector)
    connector._run()
    assert connector.view()["reports"]["koeien"]["status"] == "ok"
    assert not (connector.folder / "laatste-download-koeien.xml").exists()


def test_fertility_reports_are_not_read_when_the_herd_list_is_refused(connector):
    connector.save({**LOGIN, **REPORTS})

    def fetch(state, out):
        herd(out, 0)
        sheet_file(connector.folder / "download-koeien.xml", COW_REPORT)

    connector.fetch = fetch
    connector._run()

    assert connector.view()["status"] == "failed"
    assert not list(connector.folder.glob("download-*"))
    assert connector.view()["reports"]["koeien"]["status"] == "never"
